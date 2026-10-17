# excel → excel: два генератора под два формата.
#
# xlsx шаблон → xlsx выход: клонируем строки внутри одного openpyxl-воркбука —
#   dst._style = src._style гарантирует точное совпадение стилей.
#
# xls шаблон → xls выход: xlrd читает с formatting_info=True, xlwt пишет.
#   оба используют BIFF-константы (цвета, границы, выравнивание — те же числа),
#   поэтому перенос стилей бесшовный без каких-либо маппингов.
#
# синтаксис ячеек:
#   новый:  {{НП}}, {{Адрес объекта}}, д. {{Дом}}
#   старый: {{НП}} + ", " + {{Адрес объекта}} + ", д. " + {{Дом}}

import copy as _copy
import re
import time
from pathlib import Path

import xlrd
import openpyxl
from openpyxl.cell.rich_text import CellRichText
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from .load import traffic_light

_TAG_RE = re.compile(r"\{\{([^{}]+)\}\}")
_OLD_CONCAT_RE = re.compile(r"\}\}\s*\+|\"[^\"]*\"\s*\+\s*\{\{")
_TAG_NUM = ":num"  # {{:num}} → порядковый номер записи
_SPECS = ("addr", "time")  # {{Поле:addr}} — только по явному указанию в шаблоне


# ---------------------------------------------------------------------------
# роутер
# ---------------------------------------------------------------------------

class ExcelGenerator:
    """выбирает правильный генератор по расширению шаблона"""

    def generate(self, data_list: list, template_path, output_path, config_path=None) -> Path:
        template_path = Path(template_path)
        if template_path.suffix.lower() == ".xls":
            return XlsGenerator().generate(data_list, template_path, output_path, config_path)
        return XlsxGenerator().generate(data_list, template_path, output_path, config_path)


# ---------------------------------------------------------------------------
# xlsx → xlsx
# ---------------------------------------------------------------------------

class XlsxGenerator:

    def generate(self, data_list: list, template_path, output_path, config_path=None) -> Path:
        return self._from_xlsx(data_list, Path(template_path), output_path, config_path)

    def _from_xlsx(self, data_list: list, template_path: Path, output_path, config_path=None) -> Path:
        # rich_text=True — иначе разный кегль и цвет внутри одной ячейки
        # схлопываются в плоскую строку ещё при чтении
        wb = openpyxl.load_workbook(str(template_path), rich_text=True)

        for sheet_name in list(wb.sheetnames):
            ws_src = wb[sheet_name]
            info = _find_template_rows_xlsx(ws_src)
            if info is None:
                continue  # лист без тегов остаётся как в шаблоне

            header_count, tmpl_rows = info
            ws_out = wb.create_sheet(f"__tmp__{sheet_name}")

            for rx in range(1, header_count + 1):
                _clone_row_xlsx(ws_src, rx, ws_out, rx)

            # куда уехала каждая исходная строка: шапка на место, шаблон размножился
            row_map = {rx: [rx] for rx in range(1, header_count + 1)}

            out_row = header_count + 1
            for rec_idx, record in enumerate(data_list):
                _eval = lambda s, r, n=rec_idx + 1, **kw: _evaluate(s, r, n, config_path, **kw)
                for tmpl_rx in tmpl_rows:
                    _clone_row_xlsx(ws_src, tmpl_rx, ws_out, out_row, record, _eval)
                    row_map.setdefault(tmpl_rx, []).append(out_row)
                    out_row += 1

            _clone_merges_xlsx(ws_src, ws_out, row_map)
            _clone_cond_formats_xlsx(ws_src, ws_out, row_map)

            for col, dim in ws_src.column_dimensions.items():
                new_dim = _copy.copy(dim)
                new_dim.sheet = ws_out
                ws_out.column_dimensions[col] = new_dim

            position = wb.sheetnames.index(sheet_name)
            del wb[sheet_name]
            ws_out.title = sheet_name
            wb.move_sheet(ws_out, offset=position - wb.index(ws_out))

        out_path = _unique_path(Path(output_path).with_suffix(".xlsx"))
        wb.save(str(out_path))
        return out_path


def _find_template_rows_xlsx(ws):
    tmpl_rows = [
        row[0].row for row in ws.iter_rows()
        if any(cell.value and _TAG_RE.search(str(cell.value)) for cell in row)
    ]
    if not tmpl_rows:
        return None
    return tmpl_rows[0] - 1, tmpl_rows


def _clone_merges_xlsx(ws_src, ws_dst, row_map):
    """Переносит объединения ячеек по карте строк: шапку один раз, шаблон на каждую запись."""
    for rng in ws_src.merged_cells.ranges:
        height = rng.max_row - rng.min_row
        for start in row_map.get(rng.min_row, []):
            ws_dst.merge_cells(
                start_row=start,
                start_column=rng.min_col,
                end_row=start + height,
                end_column=rng.max_col,
            )


def _clone_cond_formats_xlsx(ws_src, ws_dst, row_map):
    """Переносит условное форматирование по той же карте строк, что и объединения."""
    for cf in ws_src.conditional_formatting:
        rows = []
        for rng in cf.sqref.ranges:
            targets = sorted(t for r, ts in row_map.items() if rng.min_row <= r <= rng.max_row for t in ts)
            # подряд идущие строки склеиваем в один диапазон
            for t in targets:
                if rows and rows[-1][0] is rng and rows[-1][2] == t - 1:
                    rows[-1][2] = t
                else:
                    rows.append([rng, t, t])
        if not rows:
            continue
        sqref = " ".join(f"{get_column_letter(rng.min_col)}{a}:{get_column_letter(rng.max_col)}{b}"
                         for rng, a, b in rows)
        for rule in cf.rules:
            ws_dst.conditional_formatting.add(sqref, rule)


def _clone_row_xlsx(ws_src, src_rx, ws_dst, dst_rx, record=None, evaluate_fn=None):
    for cell in ws_src[src_rx]:
        dst = ws_dst.cell(row=dst_rx, column=cell.column)
        value = cell.value
        # тег внутри разноформатной ячейки: подставляем как обычный текст,
        # иначе тег не найдётся и в результат уедет сам {{...}}
        if isinstance(value, CellRichText) and "{{" in str(value):
            value = str(value)
        if (record is not None and evaluate_fn is not None
                and isinstance(value, str) and "{{" in value):
            value = evaluate_fn(value, record, formula=cell.data_type == "f") or None
        # формулу протягиваем, как мышкой: относительные ссылки едут за строкой, $ стоят
        if cell.data_type == "f" and isinstance(value, str) and src_rx != dst_rx:
            value = Translator(value, origin=cell.coordinate).translate_formula(dst.coordinate)
        dst.value = value
        if cell.data_type == "s" and isinstance(value, str):
            dst.data_type = "s"  # текст остаётся текстом, даже если начинается с "="
        dst._style = cell._style  # style id валиден — один воркбук

    if src_rx in ws_src.row_dimensions:
        ws_dst.row_dimensions[dst_rx].height = ws_src.row_dimensions[src_rx].height


# ---------------------------------------------------------------------------
# xls → xls
# ---------------------------------------------------------------------------

class XlsGenerator:
    """xls → xls через xlutils.copy: копируем воркбук целиком (XF таблица, палитра,
    все стили), добавляем выходной лист В ТОМ ЖЕ воркбуке и пишем в него ячейки
    через put_cell с теми же xf_idx — никакой ручной трансляции стилей."""

    def generate(self, data_list: list, template_path, output_path, config_path=None) -> Path:
        from xlutils.copy import copy as xl_copy

        template_path = Path(template_path)
        rb = xlrd.open_workbook(str(template_path), formatting_info=True)
        wb = xl_copy(rb)  # xlwt-воркбук со всеми стилями из шаблона

        # патчим add_style: когда приходит _RawXf — возвращаем готовый idx без
        # пересчёта через весь XFStyle machinery (он для нашего объекта всё равно упадёт)
        _orig_add_style = wb.add_style
        def _patched_add_style(style):
            if isinstance(style, _RawXf):
                return style._idx
            return _orig_add_style(style)
        wb.add_style = _patched_add_style

        output_sheets = []

        for si in range(rb.nsheets):
            rs = rb.sheet_by_index(si)
            info = _find_template_rows_xls(rs)
            if info is None:
                continue

            ws_src = wb.get_sheet(si)  # xlwt-копия шаблонного листа
            header_count, tmpl_rows = info

            ws_out = wb.add_sheet(f"__out__{rs.name}")
            output_sheets.append((si, ws_out, rs.name))

            out_row = 0
            for rx in range(header_count):
                _dup_row(rs, ws_src, rx, ws_out, out_row)
                out_row += 1

            for rec_idx, record in enumerate(data_list):
                for tmpl_rx in tmpl_rows:
                    _dup_row(rs, ws_src, tmpl_rx, ws_out, out_row, record, rec_idx + 1, config_path)
                    out_row += 1

            for cx, colinfo in rs.colinfo_map.items():
                ws_out.col(cx).width = colinfo.width

            for (rlo, rhi, clo, chi) in rs.merged_cells:
                if rlo < header_count:
                    ws_out.merge(rlo, rhi - 1, clo, chi - 1)

            tmpl_row_count = len(tmpl_rows)
            tmpl_row_map = {rx: i for i, rx in enumerate(tmpl_rows)}
            for (rlo, rhi, clo, chi) in rs.merged_cells:
                if rlo in tmpl_row_map:
                    local_idx = tmpl_row_map[rlo]
                    row_span = rhi - rlo
                    for rec_idx in range(len(data_list)):
                        offset = header_count + rec_idx * tmpl_row_count + local_idx
                        ws_out.merge(offset, offset + row_span - 1, clo, chi - 1)


        # шаблонный лист меняем на выходной на том же месте, лист без тегов остаётся как в шаблоне
        replace = {si: ws_out for si, ws_out, _ in output_sheets}
        sheets = wb._Workbook__worksheets
        sheets[:] = [replace.get(si, ws) for si, ws in enumerate(sheets[:rb.nsheets])]
        for _, ws_out, name in output_sheets:
            ws_out.name = name

        out_path = _unique_path(Path(output_path).with_suffix(".xls"))
        wb.save(str(out_path))
        return out_path


class _RawXf:
    """минимальный стиль-обёртка вокруг готового xf_idx — чтобы не гонять
    xlwt через весь XFStyle machinery когда индекс уже известен."""
    class _Font:
        height = 200  # нужно только для __adjust_height; высоты ставим сами
    font = _Font()

    def __init__(self, idx: int):
        self._idx = idx

    def xf_index(self, wb):
        return self._idx


def _dup_row(rs_xlrd, ws_xlwt_src, src_rx, ws_out, dst_rx, record=None, row_num=None, config_path=None):
    """дублирует строку: xf_idx берём из xlwt-источника (тот же воркбук),
    значения — из xlrd; подстановку тегов делаем только для строковых ячеек."""
    if src_rx in rs_xlrd.rowinfo_map:
        ri = rs_xlrd.rowinfo_map[src_rx]
        ws_out.row(dst_rx).height = ri.height
        ws_out.row(dst_rx).height_mismatch = True

    src_xlwt_row = ws_xlwt_src._Worksheet__rows.get(src_rx)
    out_row_obj = ws_out.row(dst_rx)

    for cx in range(rs_xlrd.ncols):
        cell = rs_xlrd.cell(src_rx, cx)

        xf_idx = 0
        if src_xlwt_row:
            xlwt_cell = src_xlwt_row._Row__cells.get(cx)
            if xlwt_cell:
                xf_idx = xlwt_cell.xf_idx

        style = _RawXf(xf_idx)

        val = cell.value
        if record is not None and isinstance(val, str) and "{{" in val:
            val = _evaluate(val, record, row_num, config_path)

        ctype = cell.ctype
        if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            out_row_obj.set_cell_blank(cx, style)
        elif ctype == xlrd.XL_CELL_TEXT:
            out_row_obj.set_cell_text(cx, val if val is not None else "", style)
        elif ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_DATE):
            out_row_obj.set_cell_number(cx, val, style)
        elif ctype == xlrd.XL_CELL_BOOLEAN:
            out_row_obj.set_cell_boolean(cx, val, style)
        else:
            out_row_obj.set_cell_blank(cx, style)


def _find_template_rows_xls(rs):
    tmpl_rows = [
        rx for rx in range(rs.nrows)
        if any(
            isinstance(rs.cell_value(rx, cx), str) and "{{" in rs.cell_value(rx, cx)
            for cx in range(rs.ncols)
        )
    ]
    if not tmpl_rows:
        return None
    return tmpl_rows[0], tmpl_rows


# ---------------------------------------------------------------------------
# общее
# ---------------------------------------------------------------------------

def _evaluate(template_str: str, record: dict, row_num: int | None = None, config_path=None, formula=False):
    if not isinstance(template_str, str) or not template_str or "{{" not in template_str:
        return template_str

    def _value(tag: str) -> str:
        tag = tag.strip()
        if row_num is not None and tag == _TAG_NUM:
            return str(row_num)
        key, sep, spec = tag.rpartition(":")
        if sep and spec in _SPECS:
            return str(traffic_light(record.get(key, ""), spec, config_path=config_path))
        return str(record.get(tag, ""))

    def _replace(tag: str) -> str:
        # в формуле кавычка внутри строки Excel пишется двойной: оф."А" -> оф.""А""
        return _value(tag).replace('"', '""') if formula else _value(tag)

    # старая склейка через + бывает только в тексте; в формуле + это сложение Excel
    if not formula and _OLD_CONCAT_RE.search(template_str):
        parts = re.findall(r'\{\{([^{}]+)\}\}|"([^"]*)"', template_str)
        return "".join(_replace(f) if f else lit for f, lit in parts)
    return _TAG_RE.sub(lambda m: _replace(m.group(1)), template_str)



def _unique_path(path: Path) -> Path:
    if path.exists():
        return path.with_stem(f"{path.stem}_{int(time.time())}")
    return path


